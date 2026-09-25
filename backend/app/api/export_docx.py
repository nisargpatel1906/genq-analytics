import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from fastapi.responses import StreamingResponse

from app.api.chart_builder import build_charts

def _set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def _add_heading_styled(doc, text, level, accent_color=RGBColor(0x1A, 0x56, 0xDB)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(text)
    run.font.name = 'Georgia' if level == 1 else 'Arial'
    if level == 1:
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = accent_color
    else:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    return p

def _render_dynamic_sections(doc, ai_report, stats, charts_list, c_accent_rgb):
    """Renders dynamic reportSections from the AI report.
    Returns True if dynamic sections were rendered, False to fall back to legacy."""
    report_sections = ai_report.get("reportSections", [])
    if not report_sections or not isinstance(report_sections, list):
        return False

    section_num = 3  # Start at 3 (after Methodology and Executive Summary)

    for section in report_sections:
        sec_type = section.get("type", "narrative")
        sec_title = section.get("title", "Analysis")

        _add_heading_styled(doc, f"{section_num}. {sec_title}", 1)

        if sec_type == "data_overview":
            content = section.get("content", "")
            if content:
                doc.add_paragraph(content)
            doc.add_paragraph()

        elif sec_type == "findings_group":
            narrative = section.get("narrative", "")
            if narrative:
                doc.add_paragraph(narrative)

            findings = section.get("findings", [])
            for i, f in enumerate(findings, 1):
                title_t = f.get("title") or f"Finding {i}"
                detail = f.get("detail") or ""
                conf = f.get("confidence") or 0
                effect = f.get("effect_size") or ""
                practical = f.get("practical_significance") or ""

                _add_heading_styled(doc, f"{i}. {title_t}", 2, c_accent_rgb)
                
                if detail:
                    doc.add_paragraph(detail)
                
                if effect:
                    p_e = doc.add_paragraph()
                    p_e.paragraph_format.left_indent = Inches(0.25)
                    run_el = p_e.add_run("Effect Size: ")
                    run_el.font.bold = True
                    p_e.add_run(effect)
                
                if practical:
                    p_p = doc.add_paragraph()
                    p_p.paragraph_format.left_indent = Inches(0.25)
                    run_pl = p_p.add_run("Practical Significance: ")
                    run_pl.font.italic = True
                    p_p.add_run(practical)
                
                if conf:
                    p_c = doc.add_paragraph()
                    p_c.paragraph_format.left_indent = Inches(0.25)
                    p_c.paragraph_format.space_after = Pt(12)
                    run_c = p_c.add_run(f"AI confidence: {conf}%")
                    run_c.font.size = Pt(8.5)
                    run_c.font.italic = True
                    run_c.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            
            doc.add_paragraph()

        elif sec_type == "trend_analysis":
            content = section.get("content", "")
            if content:
                doc.add_paragraph(content)
            doc.add_paragraph()

        elif sec_type == "data_table":
            headers = section.get("headers", [])
            rows = section.get("rows", [])
            if headers and rows:
                table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                table.autofit = False
                
                # Calculate equal widths (total ~ 6 inches)
                col_width = Inches(6.0 / len(headers))

                for j, h in enumerate(headers):
                    cell = table.cell(0, j)
                    cell.width = col_width
                    cell.text = h
                    _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
                    _set_cell_background(cell, "111111")
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(0)
                        for r in p.runs:
                            r.font.bold = True
                            r.font.size = Pt(9)
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

                for i, row in enumerate(rows, 1):
                    for j, val in enumerate(row):
                        cell = table.cell(i, j)
                        cell.width = col_width
                        cell.text = str(val)
                        _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
                        if i % 2 == 1:
                            _set_cell_background(cell, "F9FAFB")
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(0)
                            for r in p.runs:
                                r.font.size = Pt(9)
                                r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

                doc.add_paragraph()

        elif sec_type == "comparison":
            content = section.get("content", "")
            if content:
                doc.add_paragraph(content)
            doc.add_paragraph()

        elif sec_type == "anomalies":
            anomalies = section.get("anomalies", [])
            for a in anomalies:
                sev = str(a.get("severity", "medium")).upper()
                desc = a.get("description", "")
                impact = a.get("businessImpact", "")

                p_ah = doc.add_paragraph()
                p_ah.paragraph_format.keep_with_next = True
                run_ah = p_ah.add_run(f"{a.get('column', 'Unknown')} [{sev}]")
                run_ah.font.bold = True
                run_ah.font.size = Pt(11)

                if desc:
                    p_desc = doc.add_paragraph(desc)
                    p_desc.paragraph_format.left_indent = Inches(0.25)
                if impact:
                    p_i = doc.add_paragraph()
                    p_i.paragraph_format.left_indent = Inches(0.25)
                    p_i.paragraph_format.space_after = Pt(10)
                    run_i = p_i.add_run(f"Business impact: {impact}")
                    run_i.font.size = Pt(9)
                    run_i.font.color.rgb = c_accent_rgb

            doc.add_paragraph()

        elif sec_type == "recommendations":
            recs = section.get("recommendations", [])
            for i, rec in enumerate(recs, 1):
                priority = str(rec.get("priority", "Medium")).upper()
                action = rec.get("action", "")
                rationale = rec.get("rationale", "")
                expected = rec.get("expected_outcome", "")

                p_rh = doc.add_paragraph()
                p_rh.paragraph_format.keep_with_next = True
                p_rh.paragraph_format.left_indent = Inches(0.25)

                run_num = p_rh.add_run(f"{i}. ")
                run_num.font.bold = True

                run_act = p_rh.add_run(action)
                run_act.font.bold = True

                run_pri = p_rh.add_run(f" [{priority}]")
                run_pri.font.bold = True
                run_pri.font.color.rgb = c_accent_rgb

                if rationale:
                    p_rat = doc.add_paragraph(rationale)
                    p_rat.paragraph_format.left_indent = Inches(0.5)
                
                if expected:
                    p_exp = doc.add_paragraph()
                    p_exp.paragraph_format.left_indent = Inches(0.5)
                    run_exp = p_exp.add_run(f"Expected outcome: {expected}")
                    run_exp.font.italic = True
                    run_exp.font.color.rgb = c_accent_rgb

            doc.add_paragraph()

        elif sec_type == "narrative":
            content = section.get("content", "")
            if content:
                doc.add_paragraph(content)
            doc.add_paragraph()

        section_num += 1

    return True


def _render_relational_schema_docx(doc, manifest: dict, c_accent_rgb):
    """Renders multi-table schema lineage into the Word document."""
    if not manifest:
        return
    tables = manifest.get("tables", [])
    sql_query = manifest.get("sql_query", "")
    row_count = manifest.get("row_count", 0)
    if not tables and not sql_query:
        return

    _add_heading_styled(doc, "Multi-Source Relational Architecture & Lineage", 1)
    p_sum = doc.add_paragraph()
    p_sum.add_run(
        f"Relational federation unified {len(tables)} table(s) into an analytical dataset of "
        f"{row_count:,} records via autonomous key detection."
    )

    if tables:
        t_rows = [["Entity / Table", "Record Count", "Candidate Keys"]]
        for tbl in tables:
            t_name = str(tbl.get("table_name", tbl.get("name", "Table")))
            cnt = tbl.get("row_count", tbl.get("rows", "-"))
            cnt_str = f"{cnt:,}" if isinstance(cnt, (int, float)) else str(cnt)
            raw_keys = tbl.get("candidate_keys", tbl.get("keys", []))
            t_keys = ", ".join(raw_keys) if raw_keys else "Inferred by Schema Linker"
            t_rows.append([t_name, cnt_str, t_keys])

        table = doc.add_table(rows=len(t_rows), cols=3)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                cell.text = t_rows[i][j]
                _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        doc.add_paragraph()

    if sql_query:
        _add_heading_styled(doc, "Materialized Relational SQL Execution", 2, c_accent_rgb)
        p_sql = doc.add_paragraph()
        run_sql = p_sql.add_run(sql_query)
        run_sql.font.name = "Consolas"
        run_sql.font.size = Pt(8.5)
        doc.add_paragraph()


def _render_experimentation_docx(doc, exp_results: dict, c_accent_rgb):
    """Renders A/B testing evaluation and rollout decision into the Word document."""
    if not exp_results or not exp_results.get("is_experiment"):
        return

    _add_heading_styled(doc, "A/B Testing & Controlled Experimentation Analysis", 1)
    decision = exp_results.get("rollout_decision", "INCONCLUSIVE")
    rec = exp_results.get("recommendation", "")
    target_metric = exp_results.get("primary_metric", "Target Metric")

    p_dec = doc.add_paragraph()
    r_dec = p_dec.add_run(f"ROLLOUT DECISION: {decision}\n")
    r_dec.font.bold = True
    r_dec.font.size = Pt(11)
    p_dec.add_run(f"Hypothesis Evaluation: {rec}")

    srm = exp_results.get("srm_check", {})
    srm_passed = srm.get("passed", True)
    srm_status = "PASSED (No Traffic Allocation Bias)" if srm_passed else "WARNING: SRM VIOLATION DETECTED"
    p_srm = doc.add_paragraph(f"• Sample Ratio Mismatch (SRM) Integrity: {srm_status} (Chi-Square p={srm.get('p_value', 1.0):.4f})")
    p_srm.paragraph_format.left_indent = Inches(0.25)

    lift = exp_results.get("lift_analysis", {})
    if lift:
        c_mean = lift.get("control_mean", 0.0)
        t_mean = lift.get("treatment_mean", 0.0)
        rel_lift = lift.get("relative_lift_pct", 0.0)
        p_val = lift.get("p_value", 1.0)
        ci = lift.get("confidence_interval_95", [0.0, 0.0])
        sig = "Yes (p < 0.05)" if lift.get("statistically_significant") else "No (p >= 0.05)"

        t_rows = [
            ["Metric", "Control Mean", "Variant Mean", "Lift %", "95% CI", "Stat. Sig."],
            [target_metric, f"{c_mean:.4f}", f"{t_mean:.4f}", f"{rel_lift:+.2f}%", f"[{ci[0]:.4f}, {ci[1]:.4f}]", sig]
        ]
        table = doc.add_table(rows=2, cols=6)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                cell.text = t_rows[i][j]
                _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        doc.add_paragraph()


def generate_docx_response(report_id: str, report_data: dict) -> StreamingResponse:
    ai_report   = report_data.get("report", {})
    stats       = report_data.get("stats", {})
    filename    = report_data.get("filename", "dataset")

    doc = Document()

    # Margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Styling helper variables
    c_accent_rgb = RGBColor(0x1A, 0x56, 0xDB)
    c_dark_hex = "111111"
    c_zebra_hex = "F9FAFB"
    c_alert_hex = "FFF7ED"

    # Configure default text font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.space_after = Pt(8)

    # ── Cover Page ─────────────────────────────────────────────────────────────
    for _ in range(6):
        doc.add_paragraph()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("AI Data Analysis Report")
    run_title.font.name = 'Georgia'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run(filename)
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    if ai_report.get("domain"):
        p_dom = doc.add_paragraph()
        p_dom.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_dom.paragraph_format.space_after = Pt(24)
        run_dom = p_dom.add_run(f"Domain: {ai_report['domain']}")
        run_dom.font.size = Pt(11)
        run_dom.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    # Quality Badge Metadata
    shape = stats.get("shape", {})
    dq = stats.get("data_quality", {}) or ai_report.get("_meta", {}).get("dataQuality", {})
    quality_score = dq.get("score", "?")
    quality_grade = dq.get("grade", "?")

    rows_val = shape.get("rows")
    if isinstance(rows_val, (int, float)):
        rows_str = f"{rows_val:,}"
    else:
        rows_str = str(rows_val) if rows_val is not None else "?"
    cols_val = shape.get("columns")
    cols_str = str(cols_val) if cols_val is not None else "?"

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run(
        f"{rows_str} rows  ×  {cols_str} columns  |  "
        f"Data Quality: {quality_grade} ({quality_score}/100)"
    )
    run_meta.font.size = Pt(10)
    run_meta.font.bold = True
    run_meta.font.color.rgb = c_accent_rgb

    # Footer style line
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(100)
    run_foot = p_foot.add_run("Generated by GenQ Analytics")
    run_foot.font.size = Pt(9)
    run_foot.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    doc.add_page_break()

    # ── 1. Methodology & Data Quality ───────────────────────────────────────────
    _add_heading_styled(doc, "1. Methodology & Data Quality", 1)

    if dq and dq.get("score") is not None:
        _add_heading_styled(doc, "Data Quality Assessment", 2, c_accent_rgb)

        dq_rows = [
            ["Dimension", "Score", "Out of", "Description"],
            ["Completeness", f"{dq.get('completeness', 0):.1f}", "40", "Non-null cell ratio across all columns"],
            ["Consistency", f"{dq.get('consistency', 0):.1f}", "35", "Penalises columns with >5% statistical outliers"],
            ["Structure", f"{dq.get('structure', 0):.1f}", "25", "Well-typed categorical columns with useful cardinality"],
            [f"Total  [Grade: {quality_grade}]", f"{quality_score}", "100", ""],
        ]

        table = doc.add_table(rows=len(dq_rows), cols=4)
        table.autofit = False

        for i, row in enumerate(table.rows):
            row.cells[0].width = Inches(1.8)
            row.cells[1].width = Inches(0.8)
            row.cells[2].width = Inches(0.8)
            row.cells[3].width = Inches(3.1)

            for j, cell in enumerate(row.cells):
                cell.text = dq_rows[i][j]
                _set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if i == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        elif i == len(dq_rows) - 1:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
                        else:
                            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

                if i == 0:
                     _set_cell_background(cell, c_dark_hex)
                elif i == len(dq_rows) - 1:
                     _set_cell_background(cell, "EFF6FF")
                elif i % 2 == 1:
                     _set_cell_background(cell, c_zebra_hex)

        doc.add_paragraph()

    # Methodology section
    _add_heading_styled(doc, "Analysis Methodology", 2, c_accent_rgb)
    sampling_meta = stats.get("sampling", {}) or ai_report.get("_meta", {}).get("sampling", {})
    was_sampled = sampling_meta.get("wasDownsampled", False)
    sample_size_val = sampling_meta.get("sampleSize", shape.get("rows", "?"))
    if isinstance(sample_size_val, (int, float)):
        sample_size_str = f"{sample_size_val:,}"
    else:
        sample_size_str = str(sample_size_val) if sample_size_val is not None else "?"

    full_count_val = sampling_meta.get("fullRowCount", shape.get("rows", "?"))
    if isinstance(full_count_val, (int, float)):
        full_count_str = f"{full_count_val:,}"
    else:
        full_count_str = str(full_count_val) if full_count_val is not None else "?"

    method = sampling_meta.get("method", "full")

    method_labels = {
        "full": "Full dataset — all rows used for analysis",
        "time_series": "Time-series sampling — most recent 70% + periodic samples from history",
        "stratified": "Stratified sampling — proportional sample preserving class distributions",
        "random": "Random sampling — reproducible random sample (seed=42)",
    }
    method_desc = method_labels.get(method, method)

    if was_sampled:
        methodology_text = (
            f"This dataset contained {full_count_str} rows. "
            f"To ensure timely analysis without sacrificing accuracy, a smart sample of "
            f"{sample_size_str} rows was selected using the {method_desc} strategy. "
            f"All statistics, findings, and visualizations are derived from this representative sample. "
            f"A full-data validation pass was performed to cross-check key findings against the complete dataset."
        )
    else:
        methodology_text = (
            f"This dataset contained {full_count_str} rows — within the full-analysis threshold. "
            f"All statistics, findings, and visualizations are derived from the complete dataset."
        )
    
    # Add methodology description from AI report
    ai_methodology = ai_report.get("methodology", "")
    if ai_methodology:
        methodology_text += f"\n\nStatistical Methods Used: {ai_methodology}"

    doc.add_paragraph(methodology_text)

    # Validation Warnings
    validation_warnings = ai_report.get("_meta", {}).get("sampling", {}).get("validationWarnings", [])
    if validation_warnings:
        _add_heading_styled(doc, "Full-Data Validation Warnings", 2, c_accent_rgb)
        p_warn_intro = doc.add_paragraph(
            "The following findings showed a >20% deviation between the analysis sample and the full dataset. "
            "Interpret these with additional caution:"
        )
        p_warn_intro.runs[0].font.size = Pt(9.5)

        for w in validation_warnings:
            p_w = doc.add_paragraph()
            p_w.paragraph_format.left_indent = Inches(0.25)

            run_lbl = p_w.add_run(f"{w.get('finding', 'Finding')}: ")
            run_lbl.font.bold = True
            run_lbl.font.size = Pt(9.5)

            run_body = p_w.add_run(
                f"Column '{w.get('column')}' — sample value {w.get('sample_value')}, "
                f"full-data mean {w.get('full_data_value')} ({w.get('deviation_pct')}% deviation)."
            )
            run_body.font.size = Pt(9.5)
            run_body.font.italic = True
            run_body.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)

        doc.add_paragraph()

    # ── 2. Executive Summary ────────────────────────────────────────────────────
    exec_sum = ai_report.get("executiveSummary", "")
    if exec_sum:
        _add_heading_styled(doc, "2. Executive Summary", 1)
        doc.add_paragraph(exec_sum)
        doc.add_paragraph()

    # ── Multi-Source Relational Architecture & Lineage ─────────────────────────
    relational_manifest = ai_report.get("relational_manifest", {}) or report_data.get("relational_manifest", {})
    if relational_manifest:
        _render_relational_schema_docx(doc, relational_manifest, c_accent_rgb)

    # ── Build charts for the document ───────────────────────────────────────────
    charts = build_charts(report_data)

    # ── Try dynamic sections first ──────────────────────────────────────────────
    used_dynamic = _render_dynamic_sections(doc, ai_report, stats, charts, c_accent_rgb)

    if used_dynamic:
        # Insert visualizations as a separate section after the dynamic content
        if charts:
            doc.add_page_break()
            _add_heading_styled(doc, "Data Visualizations & Interpretation", 1)
            doc.add_paragraph(
                "Each chart below was generated based on the actual statistical findings — "
                "not a generic template. An interpretation paragraph follows every chart."
            )

            for ch in charts:
                _add_heading_styled(doc, ch["title"], 2, c_accent_rgb)
                
                try:
                    ch["buf"].seek(0)
                    doc.add_picture(ch["buf"], width=Inches(5.8))
                except Exception as e:
                    p_err = doc.add_paragraph(f"[Chart Image Could Not Be Rendered: {e}]")
                    p_err.runs[0].font.color.rgb = RGBColor(0xDC, 0x26, 0x26)

                p_interp = doc.add_paragraph()
                p_interp.paragraph_format.left_indent = Inches(0.25)
                p_interp.paragraph_format.space_before = Pt(6)
                p_interp.paragraph_format.space_after = Pt(14)

                run_lbl = p_interp.add_run("Interpretation: ")
                run_lbl.font.bold = True
                run_lbl.font.size = Pt(9.5)
                run_lbl.font.color.rgb = c_accent_rgb

                run_text = p_interp.add_run(ch["interpretation"])
                run_text.font.size = Pt(9.5)
                run_text.font.italic = True
                run_text.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

                doc.add_paragraph()
        
        # Limitations section
        limitations = ai_report.get("limitations", [])
        if limitations:
            _add_heading_styled(doc, "Limitations & Caveats", 1)
            for lim in limitations:
                p_lim = doc.add_paragraph(f"• {lim}")
                p_lim.paragraph_format.left_indent = Inches(0.25)
            doc.add_paragraph()

        # A/B Testing & Controlled Experimentation
        experiment_results = ai_report.get("experiment_results", {}) or report_data.get("experiment_results", {})
        if experiment_results:
            _render_experimentation_docx(doc, experiment_results, c_accent_rgb)
            
    else:
        # ── LEGACY FALLBACK: Use the old fixed-section rendering ────────────────

        # ── 3. Dataset Statistics ────────────────────────────────────────────────────
        num_summary = stats.get("numeric_summary", {})
        missing = stats.get("missing_values", {})
        if num_summary:
            _add_heading_styled(doc, "3. Dataset Statistics", 1)

            stat_rows = [["Column", "Mean", "Std Dev", "Min", "Max", "Missing"]]
            for col, d in num_summary.items():
                stat_rows.append([
                    col[:22],
                    f"{(d.get('mean') or 0):.3f}",
                    f"{(d.get('std') or 0):.3f}",
                    f"{(d.get('min') or 0):.3f}",
                    f"{(d.get('max') or 0):.3f}",
                    str(missing.get(col, 0))
                ])

            table_s = doc.add_table(rows=len(stat_rows), cols=6)
            table_s.autofit = False

            for i, row in enumerate(table_s.rows):
                row.cells[0].width = Inches(1.9)
                row.cells[1].width = Inches(1.0)
                row.cells[2].width = Inches(1.0)
                row.cells[3].width = Inches(0.9)
                row.cells[4].width = Inches(0.9)
                row.cells[5].width = Inches(0.8)

                for j, cell in enumerate(row.cells):
                    cell.text = stat_rows[i][j]
                    _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)

                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        if j > 0:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            if i == 0:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            else:
                                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

                    if i == 0:
                        _set_cell_background(cell, c_dark_hex)
                    elif i % 2 == 0:
                        _set_cell_background(cell, c_zebra_hex)

            doc.add_paragraph()

        # ── 4. Data Visualizations & Interpretation ───────────────────────────────
        if charts:
            _add_heading_styled(doc, "4. Data Visualizations & Interpretation", 1)
            doc.add_paragraph(
                "Each chart below was chosen based on the actual structure of your dataset — "
                "not a generic template. An interpretation paragraph follows every chart "
                "explaining what the data is specifically showing."
            )

            for ch in charts:
                _add_heading_styled(doc, ch["title"], 2, c_accent_rgb)

                try:
                    ch["buf"].seek(0)
                    doc.add_picture(ch["buf"], width=Inches(5.8))
                except Exception as e:
                    p_err = doc.add_paragraph(f"[Chart Image Could Not Be Rendered: {e}]")
                    p_err.runs[0].font.color.rgb = RGBColor(0xDC, 0x26, 0x26)

                p_interp = doc.add_paragraph()
                p_interp.paragraph_format.left_indent = Inches(0.25)
                p_interp.paragraph_format.space_before = Pt(6)
                p_interp.paragraph_format.space_after = Pt(14)

                run_lbl = p_interp.add_run("Interpretation: ")
                run_lbl.font.bold = True
                run_lbl.font.size = Pt(9.5)
                run_lbl.font.color.rgb = c_accent_rgb

                run_text = p_interp.add_run(ch["interpretation"])
                run_text.font.size = Pt(9.5)
                run_text.font.italic = True
                run_text.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

                doc.add_paragraph()

        # ── 5. Key Findings ────────────────────────────────────────────────────────
        findings = ai_report.get("keyFindings", [])
        if findings:
            _add_heading_styled(doc, "5. Key Findings", 1)

            for i, f in enumerate(findings, 1):
                title_t = f.get("title") or f.get("finding") or f"Finding {i}"
                detail = f.get("detail") or f.get("description") or ""
                conf = f.get("confidenceScore") or f.get("confidence") or 0

                p_fh = doc.add_paragraph()
                p_fh.paragraph_format.keep_with_next = True
                run_fh = p_fh.add_run(f"{i}. {title_t}")
                run_fh.font.name = 'Arial'
                run_fh.font.size = Pt(11.5)
                run_fh.font.bold = True
                run_fh.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

                if detail:
                    doc.add_paragraph(detail)

                if conf:
                    p_c = doc.add_paragraph()
                    p_c.paragraph_format.space_after = Pt(12)
                    run_c = p_c.add_run(f"AI confidence: {conf}%")
                    run_c.font.size = Pt(8.5)
                    run_c.font.italic = True
                    run_c.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

            doc.add_paragraph()

        # ── 6. Anomalies Detected ──────────────────────────────────────────────────
        anomalies_ai = ai_report.get("anomalies", [])
        stat_anomalies = stats.get("statistical_anomalies", [])
        if anomalies_ai or stat_anomalies:
            _add_heading_styled(doc, "6. Anomalies Detected", 1)

            if stat_anomalies:
                _add_heading_styled(doc, "Statistically Flagged Columns (Z-score > 3σ)", 2, c_accent_rgb)

                tdata = [["Column", "Outlier Rows", "Column Mean"]]
                for a in stat_anomalies:
                    tdata.append([
                        a.get("column", ""),
                        str(a.get("outlier_count", "")),
                        f"{a.get('mean', 0):.4f}"
                    ])

                table_a = doc.add_table(rows=len(tdata), cols=3)
                table_a.autofit = False

                for i, row in enumerate(table_a.rows):
                    row.cells[0].width = Inches(2.2)
                    row.cells[1].width = Inches(1.8)
                    row.cells[2].width = Inches(2.5)

                    for j, cell in enumerate(row.cells):
                        cell.text = tdata[i][j]
                        _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)

                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(0)
                            if j > 0:
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in paragraph.runs:
                                run.font.size = Pt(9)
                                if i == 0:
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                else:
                                    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

                        if i == 0:
                            _set_cell_background(cell, c_dark_hex)
                        elif i % 2 == 1:
                            _set_cell_background(cell, c_alert_hex)

                doc.add_paragraph()

            if anomalies_ai:
                for a in anomalies_ai:
                    sev = str(a.get("severity", "medium")).upper()
                    desc = a.get("description", "")
                    impact = a.get("businessImpact", "")

                    p_ah = doc.add_paragraph()
                    p_ah.paragraph_format.keep_with_next = True
                    run_ah = p_ah.add_run(f"{a.get('column', 'Unknown')} [{sev}]")
                    run_ah.font.bold = True
                    run_ah.font.size = Pt(11)

                    if desc:
                        doc.add_paragraph(desc)
                    if impact:
                        p_i = doc.add_paragraph()
                        p_i.paragraph_format.space_after = Pt(10)
                        run_i = p_i.add_run(f"Business impact: {impact}")
                        run_i.font.size = Pt(9)
                        run_i.font.color.rgb = c_accent_rgb

                doc.add_paragraph()

        # ── 7. Recommendations ────────────────────────────────────────────────────
        recs = ai_report.get("recommendations", [])
        if recs:
            _add_heading_styled(doc, "7. Recommendations", 1)

            for i, rec in enumerate(recs, 1):
                priority = str(rec.get("priority", "Medium")).upper()
                action = rec.get("action", "")
                rationale = rec.get("rationale", "")

                p_rh = doc.add_paragraph()
                p_rh.paragraph_format.keep_with_next = True

                run_num = p_rh.add_run(f"{i}. ")
                run_num.font.bold = True

                run_act = p_rh.add_run(action)
                run_act.font.bold = True

                run_pri = p_rh.add_run(f" [{priority}]")
                run_pri.font.bold = True
                run_pri.font.color.rgb = c_accent_rgb

                if rationale:
                    doc.add_paragraph(rationale)

            doc.add_paragraph()

    # Save to BytesIO buffer
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    clean_filename = filename.replace(" ", "_").replace("/", "_").replace("\\", "_")
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={clean_filename}_Report.docx"}
    )
