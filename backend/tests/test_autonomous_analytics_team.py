# backend/tests/test_autonomous_analytics_team.py

import os
import io
import json
import pytest
import numpy as np
import pandas as pd
from unittest.mock import patch, MagicMock

from services.schema_linker import (
    infer_candidate_relationships,
    generate_deterministic_fallback_join,
    build_unified_analytical_dataset,
    connect_and_materialize_sql,
    RelationalDatabaseEngine,
)
from services.agent_graph import (
    experimentation_node,
    AnalysisGraphState,
)
from services.monitor import (
    calculate_psi,
    compute_drift_and_anomaly_alerts,
    send_webhook_alert,
)
from app.api.export_pdf import (
    _render_relational_schema,
    _render_experimentation_results,
)
from app.api.export_docx import (
    _render_relational_schema_docx,
    _render_experimentation_docx,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from docx.shared import RGBColor
from reportlab.lib import colors


# ============================================================================
# 1. Multi-Source / Relational Database & Joins Tests
# ============================================================================

def test_infer_candidate_relationships():
    """Verify that candidate foreign keys and join relationships are accurately detected."""
    df_customers = pd.DataFrame({
        "customer_id": [101, 102, 103, 104],
        "customer_name": ["Alice", "Bob", "Charlie", "David"],
        "region": ["North", "South", "East", "West"],
    })
    df_orders = pd.DataFrame({
        "order_id": [1, 2, 3, 4, 5],
        "customer_id": [101, 102, 101, 103, 999],
        "amount": [50.0, 120.0, 35.0, 200.0, 15.0],
    })

    dfs = {"customers": df_customers, "orders": df_orders}
    relationships = infer_candidate_relationships(dfs)

    assert len(relationships) >= 1
    rel = relationships[0]
    assert "customer_id" in rel["from_key"] or "customer_id" in rel["to_key"]
    assert rel["confidence"] > 0.5


def test_build_unified_analytical_dataset_sqlite():
    """Verify that multi-table datasets are loaded into SQLite in-memory and unified."""
    df_users = pd.DataFrame({
        "user_id": [1, 2, 3],
        "signup_channel": ["Organic", "Paid", "Referral"],
    })
    df_events = pd.DataFrame({
        "event_id": [10, 20, 30, 40],
        "user_id": [1, 2, 2, 3],
        "revenue": [25.0, 50.0, 75.0, 100.0],
    })

    dfs = {"users": df_users, "events": df_events}
    with patch("services.schema_linker.chat_completion", side_effect=Exception("Mock LLM offline")):
        unified_df, manifest = build_unified_analytical_dataset(dfs)

    assert not unified_df.empty
    assert any("user_id" in c for c in unified_df.columns)
    assert any("revenue" in c for c in unified_df.columns)
    assert any("signup_channel" in c for c in unified_df.columns)
    assert manifest["unified_row_count"] == len(unified_df)
    assert len(manifest["tables"]) == 2
    assert "JOIN" in manifest["sql_query"].upper() or "SELECT" in manifest["sql_query"].upper()


def test_sql_engine_materialization():
    """Verify that RelationalDatabaseEngine connects to SQLite and executes queries."""
    engine = RelationalDatabaseEngine("sqlite:///:memory:")
    sample_df = pd.DataFrame({
        "product_id": [1, 2, 3],
        "price": [10.5, 20.0, 35.5],
    })
    engine.load_dataframe("products", sample_df)

    res_df = engine.execute_query("SELECT product_id, price * 2 AS double_price FROM products WHERE price > 15")
    assert len(res_df) == 2
    assert "double_price" in res_df.columns
    assert list(res_df["double_price"]) == [40.0, 71.0]


# ============================================================================
# 2. A/B Testing & Experimentation Engine Tests
# ============================================================================

def test_experimentation_node_positive_lift():
    """Verify that experimentation_node detects variant columns, runs t-test, and recommends SHIP."""
    np.random.seed(42)
    n = 200
    control_vals = np.random.normal(loc=10.0, scale=2.0, size=n)
    treatment_vals = np.random.normal(loc=14.0, scale=2.0, size=n)

    df_exp = pd.DataFrame({
        "user_id": range(2 * n),
        "experiment_variant": ["control"] * n + ["variant_b"] * n,
        "conversion_rate": np.concatenate([control_vals, treatment_vals]),
    })

    state: AnalysisGraphState = {
        "df_cleaned": df_exp,
        "charts": [],
        "agent_progress": [],
    }

    result_state = experimentation_node(state)
    exp_res = result_state.get("experiment_results", {})

    assert exp_res.get("is_experiment") is True
    assert exp_res.get("variant_column") == "experiment_variant"
    assert "SHIP" in exp_res.get("rollout_decision")
    assert exp_res["lift_analysis"]["relative_lift_pct"] > 0
    assert exp_res["lift_analysis"]["statistically_significant"] is True
    assert exp_res["srm_check"]["passed"] is True
    assert len(result_state["charts"]) == 1
    assert "ab_test_lift" in result_state["charts"][0]["filename"]


def test_experimentation_node_negative_lift():
    """Verify that experimentation_node recommends DO NOT SHIP when treatment significantly degrades metric."""
    np.random.seed(42)
    n = 200
    control_vals = np.random.normal(loc=20.0, scale=2.0, size=n)
    treatment_vals = np.random.normal(loc=12.0, scale=2.0, size=n)

    df_exp = pd.DataFrame({
        "user_id": range(2 * n),
        "variant": ["A"] * n + ["B"] * n,
        "metric_score": np.concatenate([control_vals, treatment_vals]),
    })

    state: AnalysisGraphState = {
        "df_cleaned": df_exp,
        "charts": [],
        "agent_progress": [],
    }

    result_state = experimentation_node(state)
    exp_res = result_state.get("experiment_results", {})

    assert exp_res.get("is_experiment") is True
    assert "DO NOT SHIP" in exp_res.get("rollout_decision")
    assert exp_res["lift_analysis"]["relative_lift_pct"] < 0
    assert exp_res["lift_analysis"]["statistically_significant"] is True


def test_experimentation_node_srm_violation():
    """Verify that Sample Ratio Mismatch (SRM) is flagged when sample ratio is distorted."""
    np.random.seed(42)
    df_srm = pd.DataFrame({
        "user_id": range(1000),
        "group": ["control"] * 850 + ["treatment"] * 150,  # 85:15 distorted ratio
        "metric_val": np.random.normal(10, 2, 1000),
    })

    state: AnalysisGraphState = {
        "df_cleaned": df_srm,
        "charts": [],
        "agent_progress": [],
    }

    result_state = experimentation_node(state)
    exp_res = result_state.get("experiment_results", {})

    assert exp_res.get("is_experiment") is True
    assert exp_res["srm_check"]["passed"] is False
    assert "SRM Violation" in exp_res.get("rollout_decision")


def test_experimentation_node_non_experiment_dataset():
    """Verify that non-experiment datasets are passed through cleanly without error."""
    df_regular = pd.DataFrame({
        "date": ["2026-01-01", "2026-01-02", "2026-01-03"],
        "sales": [100, 150, 200],
    })

    state: AnalysisGraphState = {
        "df_cleaned": df_regular,
        "charts": [],
        "agent_progress": [],
    }

    result_state = experimentation_node(state)
    assert result_state.get("experiment_results", {}).get("is_experiment") is False


# ============================================================================
# 3. Continuous Monitoring & Distribution Drift Tests
# ============================================================================

def test_calculate_psi_stable():
    """Verify that identical distributions produce PSI near 0.0."""
    np.random.seed(42)
    base = pd.Series(np.random.normal(50, 10, 1000))
    curr = pd.Series(np.random.normal(50, 10, 1000))

    psi = calculate_psi(base, curr)
    assert psi < 0.10, f"Expected low PSI for stable distribution, got {psi}"


def test_calculate_psi_drifted():
    """Verify that drifted distributions produce high PSI (> 0.20)."""
    np.random.seed(42)
    base = pd.Series(np.random.normal(50, 10, 1000))
    curr = pd.Series(np.random.normal(90, 15, 1000))  # major shift

    psi = calculate_psi(base, curr)
    assert psi >= 0.20, f"Expected high PSI for drifted distribution, got {psi}"


def test_compute_drift_and_anomaly_alerts():
    """Verify end-to-end drift audit and health score calculation."""
    np.random.seed(42)
    base_df = pd.DataFrame({
        "price": np.random.normal(100, 15, 500),
        "category": np.random.choice(["Electronics", "Clothing", "Home"], size=500),
    })
    curr_df = pd.DataFrame({
        "price": np.random.normal(250, 30, 500),  # heavy drift
        "category": np.random.choice(["Electronics", "Clothing", "Home"], size=500),
    })

    report = compute_drift_and_anomaly_alerts(base_df, curr_df)
    assert report["overall_health_score"] < 80
    assert report["status"] in ["WARNING", "CRITICAL"]
    assert report["total_alerts"] >= 1
    assert "price" in report["feature_drifts"]
    assert report["feature_drifts"]["price"]["psi"] >= 0.20


def test_send_webhook_alert_format():
    """Verify that webhook dispatch builds a structured payload with alerts."""
    mock_report = {
        "status": "CRITICAL",
        "overall_health_score": 45,
        "alerts": [
            {"severity": "CRITICAL", "title": "Severe Drift on price", "detail": "PSI=0.45"}
        ],
    }

    with patch("requests.post") as mock_post:
        mock_post.return_value = MagicMock(status_code=200)
        res = send_webhook_alert("https://hooks.slack.com/services/test/alert", mock_report, "E-commerce Production")
        assert res["success"] is True
        assert mock_post.called
        args, kwargs = mock_post.call_args
        assert kwargs["json"]["status"] == "CRITICAL"
        assert kwargs["json"]["health_score"] == 45
        assert "E-commerce Production" in kwargs["json"]["text"]


# ============================================================================
# 4. Monitoring API Endpoints Tests
# ============================================================================

def test_monitor_compare_api(client):
    """Verify the /api/monitor/compare endpoint with JSON payloads."""
    headers = {"X-API-Key": "test_api_key"}
    np.random.seed(42)
    base_records = [{"val": float(x), "cat": "A"} for x in np.random.normal(10, 2, 50)]
    curr_records = [{"val": float(x), "cat": "A"} for x in np.random.normal(10, 2, 50)]

    payload = {
        "baseline_data": base_records,
        "current_data": curr_records,
        "dataset_label": "Test Batch",
    }

    resp = client.post("/api/monitor/compare", json=payload, headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert "overall_health_score" in data
    assert "status" in data
    assert data["status"] == "HEALTHY"


def test_monitor_snapshots_api(client):
    """Verify the /api/monitor/snapshots endpoint lists available snapshots."""
    headers = {"X-API-Key": "test_api_key"}
    resp = client.get("/api/monitor/snapshots", headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert "snapshots" in data


# ============================================================================
# 5. PDF & DOCX Export Rendering Tests
# ============================================================================

def test_export_pdf_render_relational_and_experiment():
    """Verify that _render_relational_schema and _render_experimentation_results build valid flowables."""
    S = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=S["Normal"], fontSize=14)
    h2 = ParagraphStyle("H2", parent=S["Normal"], fontSize=11)
    body = ParagraphStyle("Body", parent=S["Normal"], fontSize=9)
    bullet = ParagraphStyle("Bullet", parent=S["Normal"], fontSize=9)
    callout = ParagraphStyle("Callout", parent=S["Normal"], fontSize=9)

    story = []

    # 1. Test Relational Schema rendering
    mock_manifest = {
        "tables": [{"table_name": "orders", "row_count": 500, "candidate_keys": ["order_id"]}],
        "sql_query": "SELECT * FROM orders JOIN customers ON orders.customer_id = customers.id",
        "row_count": 500,
    }
    _render_relational_schema(story, mock_manifest, h1, h2, body, bullet, callout)
    assert len(story) > 0

    # 2. Test Experimentation rendering
    mock_exp = {
        "is_experiment": True,
        "primary_metric": "conversion_rate",
        "rollout_decision": "SHIP (Statistically Significant Lift)",
        "recommendation": "Roll out variant immediately.",
        "srm_check": {"passed": True, "p_value": 0.85},
        "lift_analysis": {
            "control_mean": 0.12,
            "treatment_mean": 0.18,
            "relative_lift_pct": 50.0,
            "p_value": 0.001,
            "statistically_significant": True,
            "confidence_interval_95": [0.03, 0.09],
        },
    }
    _render_experimentation_results(story, mock_exp, h1, h2, body, bullet, callout)
    assert len(story) > 4


def test_export_docx_render_relational_and_experiment():
    """Verify that _render_relational_schema_docx and _render_experimentation_docx append to Document."""
    doc = Document()
    accent = RGBColor(0x1A, 0x56, 0xDB)

    mock_manifest = {
        "tables": [{"table_name": "users", "row_count": 100}],
        "sql_query": "SELECT * FROM users",
        "row_count": 100,
    }
    _render_relational_schema_docx(doc, mock_manifest, accent)
    assert len(doc.paragraphs) > 0

    mock_exp = {
        "is_experiment": True,
        "primary_metric": "revenue",
        "rollout_decision": "SHIP",
        "recommendation": "Ship to 100% of users.",
        "srm_check": {"passed": True, "p_value": 0.99},
        "lift_analysis": {
            "control_mean": 10.0,
            "treatment_mean": 15.0,
            "relative_lift_pct": 50.0,
            "p_value": 0.0001,
            "statistically_significant": True,
            "confidence_interval_95": [3.0, 7.0],
        },
    }
    _render_experimentation_docx(doc, mock_exp, accent)
    assert len(doc.paragraphs) > 2
